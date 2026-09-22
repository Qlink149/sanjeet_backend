"""Non-destructive merge of the master contact spreadsheet into sanjeet.leads.

Unlike ``import_coachee_json`` this never deletes: every contact already in the
collection keeps its ``lead_id``, quiz state, masterclass registrations and
engagement log. The sheet only supplies curated fields (name, email, status,
interests) and fills gaps.
"""

from collections import defaultdict
from datetime import datetime, timezone
from uuid import uuid4

from qlink_chatbot.database.collections import leads
from qlink_chatbot.database.leads import canonicalize_product, map_pipeline
from qlink_chatbot.utils.logger_config import logger
from qlink_chatbot.utils.phone import (
    classify_phone,
    digits_only,
    normalize_wa_phone,
    phone_lookup_variants,
    pick_sendable_phone,
)

SHEET_NAME = "Master Database"
IMPORT_TAG = "Final_Master_Contact_Database_Updated.xlsx"

# The sheet is a strict superset of the nine other tabs, which are just views of
# it (verified: every sub-sheet phone resolves to a Master row). Import Master
# only; the "Matched Database" column already carries the sub-sheet label.
COLUMNS = [
    "name",
    "phone_raw",
    "email",
    "interest",
    "status",
    "sources",
    "segments",
    "notes",
    "wa_ready_col",  # unreliable ("phone cell non-blank") — re-derived, not used
]


def preclean_phone(raw) -> str:
    """Repair mangled prefixes that normalize_wa_phone alone cannot see through.

    Every rule below only fires on digit strings that are already unusable, so
    well-formed numbers fall through untouched.
    """
    d = digits_only(raw)
    if not d:
        return ""
    if d.startswith("00"):
        d = d[2:]
    # US international dialling prefix left on the front: 011 + 91 + number
    if d.startswith("011") and len(d) >= 13:
        d = d[3:]
    # Country code applied twice: 91 91 98204 29882
    while d.startswith("9191") and len(d) > 12:
        d = d[2:]
    # Indian country code glued onto a foreign E.164: 91 + 971 50 558 4061
    if d.startswith("91") and len(d) > 12 and d[2:].startswith("971"):
        d = d[2:]
    # 91 + national trunk 0 + 10-digit: 91 0 98999 55555
    if len(d) == 13 and d.startswith("910") and d[3] in "6789":
        d = "91" + d[3:]
    return d


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _split_multi(value) -> list[str]:
    """Sheet uses ';' between values; some cells also use ','."""
    out, seen = [], set()
    for chunk in _text(value).replace(";", ",").split(","):
        item = chunk.strip()
        if item and item not in seen:
            seen.add(item)
            out.append(item)
    return out


def parse_master_sheet(path: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"{path} has no '{SHEET_NAME}' sheet")
    ws = wb[SHEET_NAME]
    rows = []
    for i, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and _text(v) for v in values):
            continue
        row = dict(zip(COLUMNS, list(values) + [None] * len(COLUMNS)))
        row["excel_row"] = i
        rows.append(row)
    wb.close()
    return rows


def _collapse(group: list[dict]) -> dict:
    """Last sheet row wins for scalars; lists union. Mirrors _person_from_rows."""
    group = sorted(group, key=lambda r: r["excel_row"])
    latest = group[-1]

    def last(field):
        for row in reversed(group):
            if _text(row.get(field)):
                return _text(row.get(field))
        return ""

    names, products, segments, sources = [], [], [], []
    for row in group:
        name = _text(row.get("name"))
        if name and name not in names:
            names.append(name)
        for tag in canonicalize_product(_text(row.get("interest")).replace(";", ",")):
            if tag not in products:
                products.append(tag)
        for seg in _split_multi(row.get("segments")):
            if seg not in segments:
                segments.append(seg)
        for src in _split_multi(row.get("sources")):
            if src not in sources:
                sources.append(src)

    status = last("status")
    return {
        "excel_row": latest["excel_row"],
        "name": names[-1] if names else "",
        "aka": list(names[:-1]),
        "email": last("email"),
        "phone_raw": last("phone_raw"),
        "products": products,
        "product_raw": last("interest"),
        "status_raw": status,
        "pipeline": confident_pipeline(status),
        "source": "; ".join(sources),
        "segments": segments,
        "remarks": last("notes"),
    }


def confident_pipeline(status: str) -> str:
    """Pipeline from a sheet status, but only when the mapper actually matched.

    map_pipeline() returns "nurture" both for a real nurture signal and as its
    fallback for anything it does not recognise. The master sheet uses its own
    vocabulary ("Done", "3; Done", "Invited & Attending"), so a fallback here is
    absence of evidence, not evidence of nurture — treating it as a match would
    silently downgrade converted clients into the nurture audience.
    """
    if not status:
        return ""
    mapped = map_pipeline(status)
    if mapped != "nurture":
        # an explicit signal ("No Show", "Above Budget") outranks anything below
        return mapped
    # Sheet vocabulary: "Done" (often "3; Done", "5; Done" — the leading number
    # is the session count) means the engagement was completed, i.e. converted.
    # map_pipeline has no rule for it and would fall through to nurture.
    if "done" in status.lower():
        return "converted"
    return ""


def _new_doc(person: dict, stored: str, cls: str, ready: bool, now: datetime) -> dict:
    variants = phone_lookup_variants(stored) if stored else []
    return {
        "lead_id": str(uuid4()),
        "name": person["name"] or None,
        "aka": person["aka"],
        "email": person["email"] or None,
        "contact_number": stored or None,
        "contact_numbers": variants,
        "phone_class": cls,
        "whatsapp_ready": bool(stored) and ready,
        "source": person["source"] or None,
        "products": person["products"],
        "product_raw": person["product_raw"] or None,
        "pipeline": person["pipeline"] or "unknown",
        "status_raw": person["status_raw"] or None,
        "segments": person["segments"],
        "remarks": person["remarks"] or None,
        "history": [],
        "imported_from": IMPORT_TAG,
        "created_at": now,
        "updated_at": now,
    }


def _merge_ops(person: dict, doc: dict, stored: str, now: datetime):
    """Sheet wins on curated fields. Bot-maintained state is never touched.

    Returns (set_fields, add_to_set) — and never mentions quiz_archetype,
    quiz_answers, quiz_access_*, masterclass_*, engagement, history, lead_id
    or created_at.
    """
    set_fields: dict = {"updated_at": now}
    add: dict = {}

    if person["name"] and person["name"] != _text(doc.get("name")):
        set_fields["name"] = person["name"]
    for field in ("email", "status_raw", "product_raw", "remarks"):
        if person[field]:
            set_fields[field] = person[field]
    if person["pipeline"]:
        set_fields["pipeline"] = person["pipeline"]
    elif not _text(doc.get("pipeline")):
        set_fields["pipeline"] = "unknown"
    # source is bot-owned once set (e.g. "Money Ceiling Quiz") — fill only if empty
    if person["source"] and not _text(doc.get("source")):
        set_fields["source"] = person["source"]

    displaced = _text(doc.get("name"))
    aka = [a for a in person["aka"] if a and a != person["name"]]
    if displaced and "name" in set_fields and displaced != person["name"]:
        aka.append(displaced)
    if aka:
        add["aka"] = {"$each": aka}
    if person["products"]:
        add["products"] = {"$each": person["products"]}
    if person["segments"]:
        add["segments"] = {"$each": person["segments"]}
    if stored:
        add["contact_numbers"] = {"$each": phone_lookup_variants(stored)}
        # Keep the group key pinned to contact_number: if this doc's primary is
        # not sendable but the sheet number is, promote it. Otherwise
        # recompute_lead_phones() could regroup the doc onto a secondary number
        # and hard-delete whichever doc loses.
        if classify_phone(stored)[1] and not classify_phone(doc.get("contact_number"))[1]:
            cls, ready = classify_phone(stored)
            set_fields["contact_number"] = stored
            set_fields["phone_class"] = cls
            set_fields["whatsapp_ready"] = ready
    return set_fields, add


def _merge_into_pending(person: dict, doc: dict, stored: str) -> None:
    """Fold a sheet person into a doc that is still queued for insert.

    A pending doc has no lead_id in the database yet, so queueing an update
    against it would match nothing and silently drop the row.
    """
    if person["name"] and person["name"] != doc.get("name"):
        if doc.get("name"):
            doc["aka"] = doc.get("aka", []) + [doc["name"]]
        doc["name"] = person["name"]
    for field in ("email", "status_raw", "product_raw", "remarks"):
        if person[field]:
            doc[field] = person[field]
    if person["pipeline"]:
        doc["pipeline"] = person["pipeline"]
    if person["source"] and not doc.get("source"):
        doc["source"] = person["source"]
    for field, values in (
        ("products", person["products"]),
        ("segments", person["segments"]),
        ("aka", person["aka"]),
        ("contact_numbers", phone_lookup_variants(stored) if stored else []),
    ):
        current = doc.get(field) or []
        for value in values:
            if value and value not in current:
                current.append(value)
        doc[field] = current
    if stored and classify_phone(stored)[1] and not classify_phone(doc.get("contact_number"))[1]:
        cls, ready = classify_phone(stored)
        doc["contact_number"] = stored
        doc["phone_class"] = cls
        doc["whatsapp_ready"] = ready


def _projected_key(contact_number, contact_numbers) -> str:
    sendable, _, _ = pick_sendable_phone(contact_number, contact_numbers)
    return sendable


def merge_master_contacts(rows: list[dict], dry_run: bool = True) -> dict:
    """Merge parsed sheet rows into leads. Never deletes, never wipes."""
    now = datetime.now(timezone.utc)

    docs = list(leads.find({}))
    by_variant: dict = {}
    for doc in docs:
        for value in [doc.get("contact_number"), *(doc.get("contact_numbers") or [])]:
            if value:
                for variant in phone_lookup_variants(value):
                    by_variant.setdefault(variant, doc)
    by_email: dict = {}
    for doc in docs:
        email = _text(doc.get("email")).lower()
        if email:
            by_email.setdefault(email, doc)
    by_name_nophone: dict = {}
    for doc in docs:
        if not _text(doc.get("contact_number")):
            key = " ".join(_text(doc.get("name")).lower().split())
            if key:
                by_name_nophone.setdefault(key, doc)

    # --- pass 1: group sheet rows by repaired phone -------------------------
    phone_groups: dict = defaultdict(list)
    no_phone_rows: list[dict] = []
    rescued = 0
    for row in rows:
        raw = row.get("phone_raw")
        cleaned = preclean_phone(raw)
        key = normalize_wa_phone(cleaned)
        if key:
            if classify_phone(cleaned)[1] and not classify_phone(raw)[1]:
                rescued += 1
            phone_groups[key].append(row)
        else:
            no_phone_rows.append(row)

    dup_groups = sum(1 for g in phone_groups.values() if len(g) > 1)

    # --- pass 2: resolve each phone group against the DB --------------------
    updates: list = []
    inserts: list = []
    touched_ids: set = set()
    matched_phones = 0
    pending_merged = 0
    pending_by_variant: dict = {}
    # projected final state, keyed by an identity for existing docs
    projected: dict = {
        id(doc): [doc.get("contact_number"), list(doc.get("contact_numbers") or [])]
        for doc in docs
    }

    for key, group in phone_groups.items():
        person = _collapse(group)
        cleaned = preclean_phone(person["phone_raw"])
        stored, cls, ready = pick_sendable_phone(cleaned)
        stored = stored or key

        hit = None
        pending_hit = None
        for variant in phone_lookup_variants(stored):
            if variant in by_variant:
                hit = by_variant[variant]
                break
            if variant in pending_by_variant:
                pending_hit = pending_by_variant[variant]
                break

        if pending_hit is not None:
            # already queued for insert by an earlier sheet row — fold in place
            _merge_into_pending(person, pending_hit, stored)
            pending_merged += 1
            slot = projected[id(pending_hit)]
            slot[0] = pending_hit["contact_number"]
            slot[1] = list(pending_hit["contact_numbers"])
            for variant in phone_lookup_variants(stored):
                pending_by_variant.setdefault(variant, pending_hit)
        elif hit is not None:
            matched_phones += 1
            touched_ids.add(hit["lead_id"])
            set_fields, add = _merge_ops(person, hit, stored, now)
            updates.append((hit["lead_id"], set_fields, add))
            slot = projected[id(hit)]
            if "contact_number" in set_fields:
                slot[0] = set_fields["contact_number"]
            slot[1].extend(phone_lookup_variants(stored))
        else:
            doc = _new_doc(person, stored, cls, ready, now)
            inserts.append(doc)
            projected[id(doc)] = [doc["contact_number"], list(doc["contact_numbers"])]
            for variant in phone_lookup_variants(stored):
                pending_by_variant.setdefault(variant, doc)

    # --- pass 3: rows with no usable phone ----------------------------------
    nophone_by_email: dict = defaultdict(list)
    nophone_by_name: dict = defaultdict(list)
    for row in no_phone_rows:
        email = _text(row.get("email")).lower()
        if email:
            nophone_by_email[email].append(row)
        else:
            nophone_by_name[" ".join(_text(row.get("name")).lower().split())].append(row)

    nophone_merged = 0
    for email, group in nophone_by_email.items():
        person = _collapse(group)
        hit = by_email.get(email)
        if hit is not None:
            nophone_merged += 1
            touched_ids.add(hit["lead_id"])
            set_fields, add = _merge_ops(person, hit, "", now)
            updates.append((hit["lead_id"], set_fields, add))
        else:
            doc = _new_doc(person, "", "missing", False, now)
            inserts.append(doc)
            projected[id(doc)] = [None, []]
            by_email[email] = doc

    for name_key, group in nophone_by_name.items():
        if not name_key:
            continue
        person = _collapse(group)
        # only merge onto an existing doc that has no phone — avoids collapsing
        # two different people who happen to share a common name
        hit = by_name_nophone.get(name_key)
        if hit is not None:
            nophone_merged += 1
            touched_ids.add(hit["lead_id"])
            set_fields, add = _merge_ops(person, hit, "", now)
            updates.append((hit["lead_id"], set_fields, add))
        else:
            doc = _new_doc(person, "", "missing", False, now)
            inserts.append(doc)
            projected[id(doc)] = [None, []]
            by_name_nophone[name_key] = doc

    # --- invariant: no two docs may share a sendable key --------------------
    # recompute_lead_phones() runs at every boot, groups by exactly this key and
    # HARD-DELETES all but the newest doc in each group. A collision here means
    # the import would silently destroy a contact hours later, on restart.
    keys: dict = defaultdict(list)
    for slot in projected.values():
        key = _projected_key(slot[0], slot[1])
        if key:
            keys[key].append(slot)
    collisions = {k: v for k, v in keys.items() if len(v) > 1}

    report = {
        "sheet_rows": len(rows),
        "distinct_phones": len(phone_groups),
        "in_sheet_dup_groups": dup_groups,
        "rescued_by_preclean": rescued,
        "matched_existing_phones": matched_phones,
        "no_phone_rows": len(no_phone_rows),
        "no_phone_merged": nophone_merged,
        "merged_into_pending_insert": pending_merged,
        "updates": len(updates),
        "inserts": len(inserts),
        "deletes": 0,
        "existing_docs_before": len(docs),
        "existing_docs_untouched": len(docs) - len(touched_ids),
        "projected_total": len(docs) + len(inserts),
        "sendable_key_collisions": len(collisions),
        "dry_run": dry_run,
    }

    # Every queued update must target a lead_id that actually exists, otherwise
    # update_one() matches nothing and the sheet row is silently discarded.
    db_ids = {d["lead_id"] for d in docs}
    phantom = [lead_id for lead_id, _, _ in updates if lead_id not in db_ids]
    if phantom:
        raise SystemExit(
            "ABORT: %d update(s) target a lead_id not present in the database "
            "(rows would be silently dropped): %s" % (len(phantom), phantom[:10])
        )

    if collisions:
        report["collision_keys"] = list(collisions)[:20]
        raise SystemExit(
            "ABORT: %d sendable-key collision(s) — recompute_lead_phones() would "
            "hard-delete a contact on the next boot. Keys: %s"
            % (len(collisions), ", ".join(list(collisions)[:20]))
        )

    if dry_run:
        return report

    for lead_id, set_fields, add in updates:
        update: dict = {"$set": set_fields}
        if add:
            update["$addToSet"] = add
        leads.update_one({"lead_id": lead_id}, update)
    if inserts:
        leads.insert_many(inserts)

    report["actual_total"] = leads.count_documents({})
    logger.info("Merged master contact sheet", extra=report)
    return report
